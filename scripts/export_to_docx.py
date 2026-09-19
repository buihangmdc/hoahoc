"""
Script: export_to_docx.py
Purpose: Convert Markdown lesson plans and worksheets to styled Word documents (.docx).
Features:
- Math Cleaner: Converts LaTeX math syntax to clean pedagogical Unicode characters.
- Table Support: Renders Markdown tables into styled Word tables (Navy headers, borders, padding).
- Rich Formatting: Bold, italic, bullet lists, blockquotes, horizontal dividers.
- Standard Administrative Layout: A4 margins, Times New Roman, 1.2 line spacing.
"""

import os
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import latex2mathml.converter

# Load MML2OMML.XSL if available
XSLT_PATHS = [
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
]
transform = None
try:
    import lxml.etree as ET
    for p in XSLT_PATHS:
        if os.path.exists(p):
            xslt_doc = ET.parse(p)
            transform = ET.XSLT(xslt_doc)
            break
except Exception:
    transform = None

SUP_MAP = {
    '0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
    '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
    '+': '⁺', '-': '⁻', '=': '⁼', '(': '⁽', ')': '⁾',
    'n': 'ⁿ', 'i': 'ⁱ'
}
SUB_MAP = {
    '0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄',
    '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉',
    '+': '₊', '-': '₋', '=': '₌', '(': '₍', ')': '₎',
    'a': 'ₐ', 'e': 'ₑ', 'o': 'ₒ', 'x': 'ₓ', 'h': 'ₕ', 'k': 'ₖ', 'l': 'ₗ', 'm': 'ₘ', 'n': 'ₙ', 'p': 'ₚ', 's': 'ₛ', 't': 'ₜ'
}

def to_sup(s):
    return ''.join(SUP_MAP.get(c, c) for c in s)

def to_sub(s):
    return ''.join(SUB_MAP.get(c, c) for c in s)

def latex_to_omml(latex_code):
    if not transform:
        return None
    code = latex_code.strip()
    code = code.replace(r'\bar{A}', r'\overline{A}')
    code = code.replace(r'\bar{', r'\overline{')
    try:
        mathml = latex2mathml.converter.convert(code)
        tree = ET.fromstring(mathml)
        omml = transform(tree)
        return ET.tostring(omml, encoding='utf-8').decode('utf-8')
    except Exception:
        try:
            # Handle Vietnamese decimal commas like 75,77
            code_sanitized = re.sub(r'(\d+),(\d+)', r'\1{,}\2', code)
            mathml = latex2mathml.converter.convert(code_sanitized)
            tree = ET.fromstring(mathml)
            omml = transform(tree)
            return ET.tostring(omml, encoding='utf-8').decode('utf-8')
        except Exception:
            return None

def clean_latex_math(text):
    if not text:
        return text

    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1)/(\2)', text)
    text = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', text)
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\overline\{([^}]+)\}', r'¯\1', text)
    text = re.sub(r'\\bar\{([^}]+)\}', r'¯\1', text)
    text = re.sub(r'\\left\(', '(', text)
    text = re.sub(r'\\right\)', ')', text)
    text = re.sub(r'\\left\[', '[', text)
    text = re.sub(r'\\right\]', ']', text)
    
    # Generic superscripts and subscripts
    text = re.sub(r'\^\{([^}]+)\}', lambda m: to_sup(m.group(1)), text)
    text = re.sub(r'\_\{([^}]+)\}', lambda m: to_sub(m.group(1)), text)
    text = re.sub(r'\^([0-9\+\-])', lambda m: to_sup(m.group(1)), text)
    text = re.sub(r'\_([0-9\+\-])', lambda m: to_sub(m.group(1)), text)

    math_map = {
        r'\\cos': 'cos',
        r'\\sin': 'sin',
        r'\\tan': 'tan',
        r'\\cot': 'cot',
        r'\\cdot': '·',
        r'\\times': '×',
        r'\\approx': '≈',
        r'\\neq': '≠',
        r'\\le': '≤',
        r'\\leq': '≤',
        r'\\ge': '≥',
        r'\\geq': '≥',
        r'\\pm': '±',
        r'\\mp': '∓',
        r'\\circ': '°',
        r'\\infty': '∞',
        r'\\Delta': 'Δ',
        r'\\rightarrow': '→',
        r'\\Rightarrow': '⇒',
        r'\\Leftrightarrow': '⇔',
        r'\\dots': '...',
        r'\\ldots': '...',
        r'\\alpha_0': 'α₀',
        r'\\alpha': 'α',
        r'\\beta': 'β',
        r'\\gamma': 'γ',
        r'\\delta': 'δ',
        r'\\lambda': 'λ',
        r'\\omega': 'ω',
        r'\\Omega': 'Ω',
        r'\\varphi': 'φ',
        r'\\phi': 'φ',
        r'\\Phi': 'Φ',
        r'\\pi': 'π',
        r'\\theta': 'θ',
        r'\\mu': 'μ',
        r'\\rho': 'ρ',
        r'\\sigma': 'σ',
        r'\\tau': 'τ',
    }
    
    for k, v in math_map.items():
        text = re.sub(k, v, text)
        
    text = text.replace('$$', '')
    text = text.replace('$', '')
    text = text.replace('\\', '')
    
    return text

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
    <w:tblBorders {nsdecls("w")}>
      <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
      <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_formatted_text(paragraph, text, base_font_size=13, is_bold=False, is_italic=False, color_rgb=(0,0,0)):
    tokens = re.split(r'(\$\$[^\$]+\$\$|\$[^\$]+\$|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        
        if (token.startswith('$$') and token.endswith('$$')) or (token.startswith('$') and token.endswith('$')):
            latex_code = token[2:-2] if token.startswith('$$') else token[1:-1]
            omml = latex_to_omml(latex_code)
            if omml:
                try:
                    paragraph._p.append(parse_xml(omml))
                    continue
                except Exception:
                    pass
            # Fallback to Unicode formatted run
            run = paragraph.add_run(clean_latex_math(latex_code))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(*color_rgb)
            run.bold = is_bold
            run.italic = True
        elif token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            if '$' in inner:
                add_formatted_text(paragraph, inner, base_font_size=base_font_size, is_bold=True, is_italic=is_italic, color_rgb=color_rgb)
            else:
                run = paragraph.add_run(clean_latex_math(inner))
                run.font.name = 'Times New Roman'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = RGBColor(*color_rgb)
                run.bold = True
                run.italic = is_italic
        elif token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            if '$' in inner:
                add_formatted_text(paragraph, inner, base_font_size=base_font_size, is_bold=is_bold, is_italic=True, color_rgb=color_rgb)
            else:
                run = paragraph.add_run(clean_latex_math(inner))
                run.font.name = 'Times New Roman'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = RGBColor(*color_rgb)
                run.bold = is_bold
                run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.bold = True
            run.font.color.rgb = RGBColor(11, 37, 69)
        else:
            cleaned = clean_latex_math(token)
            run = paragraph.add_run(cleaned)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(*color_rgb)
            run.bold = is_bold
            run.italic = is_italic

def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def create_styled_document(input_md_path, output_docx_path):
    doc = Document()
    
    # Quy chuẩn trang A4 và căn lề: Trên 2cm, Dưới 2cm, Trái 2cm, Phải 1.5cm
    for section in doc.sections:
        section.page_width = Inches(8.267)    # A4 210mm
        section.page_height = Inches(11.692)  # A4 297mm
        section.top_margin = Inches(2.0 / 2.54)
        section.bottom_margin = Inches(2.0 / 2.54)
        section.left_margin = Inches(2.0 / 2.54)
        section.right_margin = Inches(1.5 / 2.54)

        # Đánh số trang ở Footer căn giữa
        footer = section.footer
        p_footer = footer.paragraphs[0]
        p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_footer.paragraph_format.space_before = Pt(6)
        run_footer = p_footer.add_run()
        run_footer.font.name = 'Times New Roman'
        run_footer.font.size = Pt(10)
        run_footer.font.color.rgb = RGBColor(100, 116, 139)
        add_page_number(run_footer)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.2
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if not os.path.exists(input_md_path):
        return

    with open(input_md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        text = line.strip()
        
        if not text:
            i += 1
            continue
            
        if text.startswith('|') and text.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                tline = lines[i].strip()
                if not re.match(r'^\|(\s*:?-+:?\s*\|)+$', tline):
                    table_lines.append(tline)
                i += 1
                
            if table_lines:
                rows_data = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.split('|')[1:-1]]
                    rows_data.append(cells)
                    
                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    table = doc.add_table(rows=len(rows_data), cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(table, color="94A3B8", sz="4")
                    
                    for r_idx, row in enumerate(rows_data):
                        for c_idx in range(num_cols):
                            cell_text = row[c_idx] if c_idx < len(row) else ""
                            cell = table.cell(r_idx, c_idx)
                            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
                            
                            if r_idx == 0:
                                set_cell_background(cell, "0B2545")
                                add_formatted_text(cell.paragraphs[0], cell_text, base_font_size=12, is_bold=True, color_rgb=(255, 255, 255))
                                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F8FAFC")
                                add_formatted_text(cell.paragraphs[0], cell_text, base_font_size=12, color_rgb=(30, 41, 59))
                                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                                
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_after = Pt(4)
            continue
            
        if text.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            add_formatted_text(p, text[2:].upper(), base_font_size=15, is_bold=True, color_rgb=(11, 37, 69))
            
        elif text.startswith('## '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, text[3:], base_font_size=13.5, is_bold=True, color_rgb=(19, 64, 116))
            
        elif text.startswith('### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, text[4:], base_font_size=13, is_bold=True, color_rgb=(37, 99, 235))
            
        elif text.startswith('#### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, text[5:], base_font_size=13, is_bold=True, is_italic=True, color_rgb=(30, 41, 59))

        elif text.startswith('$$') and text.endswith('$$'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_text(p, text, base_font_size=13)

        elif text.startswith('---') or text.startswith('___'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("______________________________________________________________________")
            run.font.color.rgb = RGBColor(203, 213, 225)
            run.font.size = Pt(10)

        elif text.startswith('![') and '](' in text and text.endswith(')'):
            m = re.match(r'!\[(.*?)\]\((.*?)\)', text)
            if m:
                caption = m.group(1).strip()
                img_path = m.group(2).strip()
                if os.path.exists(img_path):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(4.2))
                    if caption:
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_after = Pt(6)
                        r_cap = p_cap.add_run(caption)
                        r_cap.italic = True
                        r_cap.font.name = 'Times New Roman'
                        r_cap.font.size = Pt(10)
                        r_cap.font.color.rgb = RGBColor(100, 116, 139)

        elif text.startswith('> '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, text[2:], base_font_size=12.5, is_italic=True, color_rgb=(71, 85, 105))

        elif text.startswith('- ') or text.startswith('* '):
            raw_item = text[2:].strip()
            # Kiểm tra nếu là phương án trắc nghiệm (A., B., C., D.), ý đúng sai (a), b), c), d), Đáp số, Lời giải...
            # TUYỆT ĐỐI KHÔNG để gạch đầu dòng!
            is_option = bool(re.match(r'^([A-D]\.|\([A-D]\)|[a-d]\)|\([a-d]\)|[a-d]\.|Đáp số:|Hướng dẫn giải:|Lời giải:)', raw_item))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2 if is_option else 3)
            if is_option:
                add_formatted_text(p, raw_item, base_font_size=13)
            else:
                add_formatted_text(p, "- " + raw_item, base_font_size=13)

        elif text.startswith('+ '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, "+ " + text[2:], base_font_size=13)
            
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            # Thụt lề nhẹ cho các phương án trắc nghiệm nếu không có bullet
            if re.match(r'^([A-D]\.|\([A-D]\)|[a-d]\)|\([a-d]\)|[a-d]\.|Đáp số:)', text):
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(2)
            else:
                p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, text, base_font_size=13)

        i += 1

    os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)
    doc.save(output_docx_path)

if __name__ == "__main__":
    input_path = sys.argv[1] if len(sys.argv) > 1 else "dau-ra/giao-an/02-ke-hoach-bai-day-5512.md"
    output_path = sys.argv[2] if len(sys.argv) > 2 else "dau-ra/giao-an/02-ke-hoach-bai-day-5512.docx"
    create_styled_document(input_path, output_path)