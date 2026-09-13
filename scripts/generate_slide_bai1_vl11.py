"""
Script: generate_slide_bai1_vl11.py
Purpose: Generate ultra-high-quality, professional, animated PowerPoint presentation
for "Bài 1: Dao động điều hòa" (Vật lí 11 - Kết nối tri thức với cuộc sống).
15 Widescreen (16:9) slides with large typography:
- Header Title: 34pt (auto-fit)
- Card Title: 28pt (auto-fit)
- Body Content / Bullets / Options: 26-28pt (auto-fit)
- Dynamic vertical offset between titles and content to eliminate collisions
- High-resolution scientific diagrams, transitions, and on-click entrance animations.
"""

import os
import sys
import io
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Palette chuẩn sư phạm
COLOR_NAVY_DARK = RGBColor(11, 37, 69)       # #0B2545 - Nền Header & Slide Bìa
COLOR_NAVY_LIGHT = RGBColor(19, 64, 116)     # #134074 - Khối phụ
COLOR_ROYAL_BLUE = RGBColor(37, 99, 235)     # #2563EB - Điểm nhấn Xanh dương
COLOR_SKY_BLUE = RGBColor(141, 169, 196)     # #8DA9C4 - Màu chữ phụ trên nền tối
COLOR_ICE_BG = RGBColor(245, 247, 250)       # #F5F7FA - Nền slide dịu mắt
COLOR_WHITE = RGBColor(255, 255, 255)        # #FFFFFF - Nền Card
COLOR_AMBER = RGBColor(245, 158, 11)         # #F59E0B - Điểm nhấn Vàng Cam
COLOR_AMBER_BG = RGBColor(254, 243, 199)     # #FEF3C7 - Nền Hộp Chú ý
COLOR_AMBER_BORDER = RGBColor(251, 191, 36)  # #FBBF24
COLOR_EMERALD = RGBColor(16, 185, 129)       # #10B981 - Màu Xanh Lá Chuẩn
COLOR_EMERALD_BG = RGBColor(209, 250, 229)   # #D1FAE5 - Nền Đáp án Đúng
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)       # #1E293B - Màu chữ nội dung chính
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B - Màu chữ phụ
COLOR_CARD_BORDER = RGBColor(203, 213, 225)  # #CBD5E1 - Viền Card rõ nét
COLOR_RED_ACCENT = RGBColor(239, 68, 68)     # #EF4444 - Mâu thuẫn / Cảnh báo

ASSET_DIR = Path("dau-ra/bai-day/vat-ly-11/bai-1-dao-dong-dieu-hoa/assets_img")

def calculate_title_font_size(text, max_width_inches, base_size=34):
    length = len(text)
    avail_pt = max_width_inches * 72
    char_width_factor = 0.52
    req_pt = length * (base_size * char_width_factor)
    if req_pt <= avail_pt:
        return base_size
    for test_size in [32, 30, 28, 26, 24]:
        if length * (test_size * char_width_factor) <= avail_pt:
            return test_size
    return 24

def calculate_card_title_font_size(text, max_width_inches, base_size=28):
    length = len(text)
    avail_pt = max_width_inches * 72
    char_width_factor = 0.54
    req_pt = length * (base_size * char_width_factor)
    if req_pt <= avail_pt:
        return base_size
    for test_size in [26, 24, 22, 20]:
        if length * (test_size * char_width_factor) <= avail_pt:
            return test_size
    return 20

def add_slide_transition(slide, trans_type="fade"):
    trans_xml = parse_xml(f'<p:transition {nsdecls("p")} spd="med"><p:{trans_type}/></p:transition>')
    slide._element.append(trans_xml)

def add_click_animations(slide, shape_ids):
    if not shape_ids:
        return
    child_nodes = []
    bld_nodes = []
    c_tn_id = 3
    for sp_id in shape_ids:
        c_tn_id_set = c_tn_id + 2
        c_tn_id_anim = c_tn_id + 3
        child_xml = f'''
        <p:par {nsdecls("p")}>
          <p:cTn id="{c_tn_id}" fill="hold">
            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
            <p:childTnLst>
              <p:par>
                <p:cTn id="{c_tn_id + 1}" fill="hold">
                  <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                  <p:childTnLst>
                    <p:set>
                      <p:cBhvr>
                        <p:cTn id="{c_tn_id_set}" dur="1" fill="hold"/>
                        <p:tgtEl><p:spTgt spid="{sp_id}"/></p:tgtEl>
                        <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                      </p:cBhvr>
                      <p:to><p:strVal val="visible"/></p:to>
                    </p:set>
                    <p:animEffect transition="in" filter="fade">
                      <p:cBhvr>
                        <p:cTn id="{c_tn_id_anim}" dur="350"/>
                        <p:tgtEl><p:spTgt spid="{sp_id}"/></p:tgtEl>
                      </p:cBhvr>
                    </p:animEffect>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
            </p:childTnLst>
          </p:cTn>
        </p:par>
        '''
        child_nodes.append(child_xml)
        bld_nodes.append(f'<p:bldP {nsdecls("p")} spid="{sp_id}" grpId="0" animBg="1"/>')
        c_tn_id += 4
    
    children_combined = "".join(child_nodes)
    blds_combined = "".join(bld_nodes)
    timing_xml = parse_xml(f'''
    <p:timing {nsdecls("p")}>
      <p:tnLst>
        <p:par>
          <p:cTn id="1" dur="indefinite" restart="always" nodeType="tmRoot">
            <p:childTnLst>
              <p:seq concurrent="1" nextAc="seek">
                <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                  <p:childTnLst>
                    {children_combined}
                  </p:childTnLst>
                </p:cTn>
                <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
                <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
              </p:seq>
            </p:childTnLst>
          </p:cTn>
        </p:par>
      </p:tnLst>
      <p:bldLst>{blds_combined}</p:bldLst>
    </p:timing>
    ''')
    slide._element.append(timing_xml)

def set_slide_background(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def create_header(slide, activity_badge, main_title):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.35))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_NAVY_DARK
    top_bar.line.color.rgb = COLOR_NAVY_DARK
    
    accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.35), Inches(13.333), Inches(0.06))
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = COLOR_AMBER
    accent_line.line.color.rgb = COLOR_AMBER

    if activity_badge:
        badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.38), Inches(2.7), Inches(0.58))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = COLOR_NAVY_LIGHT
        badge_box.line.color.rgb = COLOR_AMBER
        badge_box.line.width = Pt(2)
        tf = badge_box.text_frame
        p = tf.paragraphs[0]
        p.text = activity_badge.upper()
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_AMBER
        p.font.name = 'Segoe UI'
        p.alignment = PP_ALIGN.CENTER
    
    title_left = Inches(0.6) if not activity_badge else Inches(3.5)
    title_width_in = 12.1 if not activity_badge else 9.3
    
    fitted_title_size = calculate_title_font_size(main_title, title_width_in, base_size=34)
    
    title_box = slide.shapes.add_textbox(title_left, Inches(0.2), Inches(title_width_in), Inches(0.95))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = main_title
    p_title.font.size = Pt(fitted_title_size)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE
    p_title.font.name = 'Segoe UI'

    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.1), Inches(12.133), Inches(0.35))
    tf_foot = footer_box.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "VẬT LÍ 11 — BÀI 1: DAO ĐỘNG ĐIỀU HÒA | KẾT NỐI TRI THỨC VỚI CUỘC SỐNG"
    p_foot.font.size = Pt(11)
    p_foot.font.color.rgb = COLOR_TEXT_MUTED
    p_foot.font.name = 'Segoe UI'

def add_card_with_autofit(slide, left, top, width, height, title="", bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, base_title_size=28):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(2.0)
    
    content_start_top = top + Inches(0.2)
    if title:
        width_in = width.inches - 0.4
        fitted_size = calculate_card_title_font_size(title, width_in, base_size=base_title_size)
        title_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), Inches(width_in), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(fitted_size)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY_DARK
        p.font.name = 'Segoe UI'
        content_start_top = top + Inches(0.72)
    
    return card, content_start_top

def add_bullet_list(slide, left, top, width, height, items, font_size=26, text_color=COLOR_TEXT_MAIN, space_after=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = text_color
        p.font.name = 'Segoe UI'
        p.space_after = Pt(space_after)
    return box

def safe_add_image(slide, img_name, left, top, width, height, border_color=COLOR_CARD_BORDER):
    img_path = ASSET_DIR / img_name
    if img_path.exists():
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        frame.fill.solid()
        frame.fill.fore_color.rgb = COLOR_WHITE
        frame.line.color.rgb = border_color
        frame.line.width = Pt(2.0)
        pad = Inches(0.06)
        pic = slide.shapes.add_picture(str(img_path), left + pad, top + pad, width - pad*2, height - pad*2)
        return frame, pic
    return None, None

def generate_vl11_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================
    # SLIDE 1: BÌA BÀI GIẢNG (TITLE SLIDE)
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_NAVY_DARK)
    add_slide_transition(s1, "fade")

    main_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3))
    main_card.fill.solid()
    main_card.fill.fore_color.rgb = RGBColor(15, 23, 42)
    main_card.line.color.rgb = COLOR_AMBER
    main_card.line.width = Pt(3.0)

    tag_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(0.9), Inches(8.333), Inches(0.55))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = COLOR_NAVY_LIGHT
    tag_box.line.color.rgb = COLOR_ROYAL_BLUE
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "VẬT LÍ 11 — CHƯƠNG I: DAO ĐỘNG (KẾT NỐI TRI THỨC)"
    p_tag.font.size = Pt(17)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_AMBER
    p_tag.alignment = PP_ALIGN.CENTER

    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(1.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p1 = tf_title.paragraphs[0]
    p1.text = "BÀI 1: DAO ĐỘNG ĐIỀU HÒA"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = 'Segoe UI'

    p2 = tf_title.add_paragraph()
    p2.text = "Thời lượng: 02 tiết (90 phút)  •  Chuẩn Công văn 5512"
    p2.font.size = Pt(24)
    p2.font.color.rgb = COLOR_SKY_BLUE
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = 'Segoe UI'

    div_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.0), Inches(3.55), Inches(5.333), Inches(0.06))
    div_bar.fill.solid()
    div_bar.fill.fore_color.rgb = COLOR_AMBER
    div_bar.line.color.rgb = COLOR_AMBER

    b1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(3.85), Inches(5.3), Inches(2.2))
    b1.fill.solid()
    b1.fill.fore_color.rgb = COLOR_NAVY_DARK
    b1.line.color.rgb = COLOR_ROYAL_BLUE
    b1.line.width = Pt(2.0)
    tf_b1 = b1.text_frame
    tf_b1.word_wrap = True
    p_b1_title = tf_b1.paragraphs[0]
    p_b1_title.text = "🎯 MỤC TIÊU TRỌNG TÂM"
    p_b1_title.font.size = Pt(22)
    p_b1_title.font.bold = True
    p_b1_title.font.color.rgb = COLOR_AMBER
    for it in [
        "Phương trình li độ x = A·cos(ωt + φ)",
        "Đồ thị hình sin x - t & Vòng tròn lượng giác"
    ]:
        p = tf_b1.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(20)
        p.font.color.rgb = COLOR_WHITE

    b2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(3.85), Inches(5.3), Inches(2.2))
    b2.fill.solid()
    b2.fill.fore_color.rgb = COLOR_NAVY_DARK
    b2.line.color.rgb = COLOR_EMERALD
    b2.line.width = Pt(2.0)
    tf_b2 = b2.text_frame
    tf_b2.word_wrap = True
    p_b2_title = tf_b2.paragraphs[0]
    p_b2_title.text = "⚡ ĐỔI MỚI NĂNG LỰC SỐ"
    p_b2_title.font.size = Pt(22)
    p_b2_title.font.bold = True
    p_b2_title.font.color.rgb = COLOR_EMERALD
    for it in [
        "Thực hành thí nghiệm ảo PhET (NLS 4.1.1)",
        "Mô phỏng phương trình với AI (AI-V11)"
    ]:
        p = tf_b2.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(20)
        p.font.color.rgb = COLOR_WHITE

    bot_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.45))
    tf_bot = bot_box.text_frame
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "Giáo viên: Nhóm Chuyên môn Vật lí  •  Năm học: 2026 - 2027"
    p_bot.font.size = Pt(16)
    p_bot.font.italic = True
    p_bot.font.color.rgb = COLOR_SKY_BLUE
    p_bot.alignment = PP_ALIGN.CENTER

    add_click_animations(s1, [b1.shape_id, b2.shape_id])

    # =========================================================
    # SLIDE 2: MỤC TIÊU BÀI HỌC
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_ICE_BG)
    add_slide_transition(s2, "push")
    create_header(s2, "YCCĐ", "MỤC TIÊU BÀI HỌC CẦN ĐẠT")

    c1, top1 = add_card_with_autofit(s2, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.3), "1. KIẾN THỨC CỐT LÕI", COLOR_WHITE, COLOR_ROYAL_BLUE, base_title_size=28)
    add_bullet_list(s2, Inches(0.8), top1 + Inches(0.1), Inches(5.5), Inches(4.3), [
        "Nắm chắc khái niệm: Dao động cơ & Dao động điều hòa.",
        "Viết thành thạo phương trình: x = A·cos(ωt + φ).",
        "Khai thác chính xác đồ thị li độ – thời gian (x - t)."
    ], font_size=26, space_after=14)

    c2, top2 = add_card_with_autofit(s2, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3), "2. NĂNG LỰC SỐ & ỨNG DỤNG", COLOR_WHITE, COLOR_EMERALD, base_title_size=28)
    add_bullet_list(s2, Inches(7.0), top2 + Inches(0.1), Inches(5.5), Inches(4.3), [
        "Vận dụng vòng tròn lượng giác giải bài toán thời gian.",
        "NLS 4.1.1: Đo đạc số liệu trên mô phỏng PhET.",
        "AI-V11: Tạo Prompt khảo sát pha dao động cùng AI."
    ], font_size=26, space_after=14)

    add_click_animations(s2, [c1.shape_id, c2.shape_id])

    # =========================================================
    # SLIDE 3: HOẠT ĐỘNG 1: KHỞI ĐỘNG (DAO ĐỘNG CƠ)
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_ICE_BG)
    add_slide_transition(s3, "push")
    create_header(s3, "HĐ 1 (7p)", "KHỞI ĐỘNG: DAO ĐỘNG CƠ TRONG TỰ NHIÊN")

    c_left, top_l3 = add_card_with_autofit(s3, Inches(0.6), Inches(1.6), Inches(5.9), Inches(4.3), "HIỆN TƯỢNG QUAN SÁT", COLOR_WHITE, COLOR_ROYAL_BLUE, base_title_size=28)
    add_bullet_list(s3, Inches(0.8), top_l3 + Inches(0.05), Inches(5.5), Inches(3.4), [
        "Chiếc xích đu đung đưa quanh vị trí đứng yên.",
        "Quả lắc đồng hồ chuyển động qua lại.",
        "Con lắc lò xo dãn rồi co liên tục."
    ], font_size=26, space_after=12)

    c_right, top_r3 = add_card_with_autofit(s3, Inches(6.8), Inches(1.6), Inches(5.9), Inches(4.3), "ĐẶC ĐIỂM CHUNG", COLOR_WHITE, COLOR_EMERALD, base_title_size=28)
    add_bullet_list(s3, Inches(7.0), top_r3 + Inches(0.05), Inches(5.5), Inches(3.4), [
        "Chuyển động lặp đi lặp lại nhiều lần.",
        "Xảy ra quanh một VỊ TRÍ CÂN BẰNG (VTCB).",
        "➔ Định nghĩa: Đó là DAO ĐỘNG CƠ!"
    ], font_size=26, space_after=12)

    bot_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.05), Inches(12.1), Inches(0.95))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = COLOR_AMBER_BG
    bot_card.line.color.rgb = COLOR_AMBER_BORDER
    bot_card.line.width = Pt(2.0)
    tf_take = bot_card.text_frame
    p_t = tf_take.paragraphs[0]
    p_t.text = "❓ VẤN ĐỀ: Hàm số toán học nào mô tả chính xác chuyển động này?"
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_NAVY_DARK
    p_t.alignment = PP_ALIGN.CENTER

    add_click_animations(s3, [c_left.shape_id, c_right.shape_id, bot_card.shape_id])

    # =========================================================
    # SLIDE 4: HOẠT ĐỘNG 2.1: PHƯƠNG TRÌNH DAO ĐỘNG ĐIỀU HÒA
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_ICE_BG)
    add_slide_transition(s4, "push")
    create_header(s4, "HĐ 2.1 (15p)", "PHƯƠNG TRÌNH DAO ĐỘNG ĐIỀU HÒA")

    c_def, top_def = add_card_with_autofit(s4, Inches(0.6), Inches(1.6), Inches(12.133), Inches(2.2), "1. ĐỊNH NGHĨA DAO ĐỘNG ĐIỀU HÒA", COLOR_WHITE, COLOR_ROYAL_BLUE, base_title_size=28)
    add_bullet_list(s4, Inches(0.8), top_def + Inches(0.05), Inches(11.7), Inches(1.4), [
        "Dao động điều hòa là dao động trong đó li độ của vật là một hàm cosin (hoặc sin) của thời gian.",
        "Quy luật biến thiên tuần hoàn: lặp lại trạng thái sau mỗi khoảng thời gian xác định."
    ], font_size=26, space_after=8)

    c_eq, top_eq = add_card_with_autofit(s4, Inches(0.6), Inches(4.0), Inches(12.133), Inches(3.0), "2. PHƯƠNG TRÌNH LI ĐỘ CHUẨN", COLOR_WHITE, COLOR_EMERALD, base_title_size=28)
    
    # Formula Highlight box
    f_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), top_eq + Inches(0.05), Inches(8.333), Inches(1.05))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = COLOR_NAVY_DARK
    f_box.line.color.rgb = COLOR_AMBER
    f_box.line.width = Pt(2.5)
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "x = A · cos(ωt + φ)"
    p_f.font.size = Pt(36)
    p_f.font.bold = True
    p_f.font.color.rgb = COLOR_AMBER
    p_f.alignment = PP_ALIGN.CENTER

    add_bullet_list(s4, Inches(0.8), top_eq + Inches(1.25), Inches(11.7), Inches(1.0), [
        "x: Li độ (cm/m)  •  A: Biên độ (A > 0)  •  ω: Tần số góc (rad/s)  •  φ: Pha ban đầu (rad)"
    ], font_size=26, text_color=COLOR_NAVY_DARK)

    add_click_animations(s4, [c_def.shape_id, c_eq.shape_id])

    # =========================================================
    # SLIDE 5: CÁC ĐẠI LƯỢNG ĐẶC TRƯNG
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_ICE_BG)
    add_slide_transition(s5, "push")
    create_header(s5, "Đại lượng", "CÁC ĐẠI LƯỢNG ĐẶC TRƯNG CỦA DAO ĐỘNG")

    c1, top_c1 = add_card_with_autofit(s5, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.3), "CHU KÌ, TẦN SỐ & TẦN SỐ GÓC", COLOR_WHITE, COLOR_ROYAL_BLUE, base_title_size=26)
    add_bullet_list(s5, Inches(0.8), top_c1 + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Chu kì T (s): Thời gian thực hiện 1 dao động toàn phần.",
        "Tần số f (Hz): Số dao động trong 1 giây.",
        "Công thức liên hệ cốt lõi:",
        "   ➔  ω = 2π/T = 2π·f",
        "   ➔  T = 1/f = 2π/ω"
    ], font_size=26, space_after=10)

    c2, top_c2 = add_card_with_autofit(s5, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3), "LI ĐỘ, BIÊN ĐỘ & PHA", COLOR_WHITE, COLOR_EMERALD, base_title_size=26)
    add_bullet_list(s5, Inches(7.0), top_c2 + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Li độ x: Vị trí so với VTCB (-A ≤ x ≤ +A).",
        "Biên độ A: Độ lệch cực đại (A > 0).",
        "Pha ban đầu φ: Trạng thái tại t = 0.",
        "Pha dao động (ωt + φ): Trạng thái tại thời điểm t bất kì."
    ], font_size=26, space_after=10)

    add_click_animations(s5, [c1.shape_id, c2.shape_id])

    # =========================================================
    # SLIDE 6: MÔ HÌNH CON LẮC LÒ XO DAO ĐỘNG ĐIỀU HÒA
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_ICE_BG)
    add_slide_transition(s6, "push")
    create_header(s6, "Mô hình", "MÔ HÌNH CON LẮC LÒ XO DAO ĐỘNG")

    safe_add_image(s6, "con_lac_lo_xo.png", Inches(0.6), Inches(1.6), Inches(12.133), Inches(5.3))

    # =========================================================
    # SLIDE 7: HĐ 2.2: VÒNG TRÒN LƯỢNG GIÁC LIÊN HỆ
    # =========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, COLOR_ICE_BG)
    add_slide_transition(s7, "push")
    create_header(s7, "HĐ 2.2 (15p)", "VÒNG TRÒN LƯỢNG GIÁC & C/Đ TRÒN ĐỀU")

    safe_add_image(s7, "vong_tron_luong_giac.png", Inches(0.6), Inches(1.6), Inches(12.133), Inches(5.3))

    # =========================================================
    # SLIDE 8: QUY TẮC XÁC ĐỊNH PHA & CHIỀU CHUYỂN ĐỘNG
    # =========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, COLOR_ICE_BG)
    add_slide_transition(s8, "push")
    create_header(s8, "Quy tắc Pha", "XÁC ĐỊNH PHA & CHIỀU CHUYỂN ĐỘNG")

    c_top, top_t8 = add_card_with_autofit(s8, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.3), "NỬA TRÊN VÒNG TRÒN (y > 0)", COLOR_WHITE, COLOR_RED_ACCENT, base_title_size=26)
    add_bullet_list(s8, Inches(0.8), top_t8 + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Góc pha dương: 0 < φ < π.",
        "Hình chiếu P chuyển động theo CHIỀU ÂM (v < 0).",
        "Vật đang đi về phía biên âm -A."
    ], font_size=26, space_after=14)

    c_bot, top_b8 = add_card_with_autofit(s8, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3), "NỬA DƯỚI VÒNG TRÒN (y < 0)", COLOR_WHITE, COLOR_EMERALD, base_title_size=26)
    add_bullet_list(s8, Inches(7.0), top_b8 + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Góc pha âm: -π < φ < 0.",
        "Hình chiếu P chuyển động theo CHIỀU DƯƠNG (v > 0).",
        "Vật đang đi về phía biên dương +A."
    ], font_size=26, space_after=14)

    add_click_animations(s8, [c_top.shape_id, c_bot.shape_id])

    # =========================================================
    # SLIDE 9: HĐ 2.3: ĐỒ THỊ LI ĐỘ – THỜI GIAN (x - t)
    # =========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, COLOR_ICE_BG)
    add_slide_transition(s9, "push")
    create_header(s9, "HĐ 2.3 (13p)", "ĐỒ THỊ LI ĐỘ – THỜI GIAN (x - t)")

    safe_add_image(s9, "do_thi_dao_dong.png", Inches(0.6), Inches(1.6), Inches(12.133), Inches(5.3))

    # =========================================================
    # SLIDE 10: DAO ĐỘNG CỦA CON LẮC ĐƠN
    # =========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, COLOR_ICE_BG)
    add_slide_transition(s10, "push")
    create_header(s10, "Con lắc đơn", "DAO ĐỘNG CỦA CON LẮC ĐƠN")

    safe_add_image(s10, "con_lac_don.png", Inches(0.6), Inches(1.6), Inches(12.133), Inches(5.3))

    # =========================================================
    # SLIDE 11: BẢNG TỔNG HỢP TRẠNG THÁI DAO ĐỘNG
    # =========================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, COLOR_ICE_BG)
    add_slide_transition(s11, "push")
    create_header(s11, "Tổng hợp", "BẢNG TRẠNG THÁI VỊ TRÍ ĐẶC BIỆT")

    safe_add_image(s11, "bang_trang_thai.png", Inches(0.6), Inches(1.6), Inches(12.133), Inches(5.3))

    # =========================================================
    # SLIDE 12: HĐ 2.4: THỰC HÀNH MÔ PHỎNG PHET & AI-V11
    # =========================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, COLOR_ICE_BG)
    add_slide_transition(s12, "push")
    create_header(s12, "HĐ 2.4 (12p)", "MÔ PHỎNG PHET & KHÁM PHÁ AI")

    c_phet, top_ph = add_card_with_autofit(s12, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.3), "THÍ NGHIỆM ẢO PHET (NLS 4.1.1)", COLOR_WHITE, COLOR_ROYAL_BLUE, base_title_size=24)
    add_bullet_list(s12, Inches(0.8), top_ph + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Truy cập mô phỏng PhET: Masses & Springs.",
        "Khảo sát: Thay đổi khối lượng m và độ cứng k.",
        "Đo chu kì dao động T bằng đồng hồ bấm giây ảo.",
        "Rút ra kết luận: Chu kì T không phụ thuộc vào biên độ A nhỏ!"
    ], font_size=26, space_after=10)

    c_ai, top_ai = add_card_with_autofit(s12, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3), "TRẢI NGHIỆM AI PROMPT (AI-V11)", COLOR_WHITE, COLOR_EMERALD, base_title_size=24)
    add_bullet_list(s12, Inches(7.0), top_ai + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Câu lệnh mẫu Prompting:",
        '  💬 "Hãy giải thích ý nghĩa hình học của pha ban đầu φ và mô phỏng đồ thị khi φ = π/2 và φ = -π/2."',
        "Đối chiếu kết quả phân tích cùng AI với đồ thị thực nghiệm SGK."
    ], font_size=26, space_after=10)

    add_click_animations(s12, [c_phet.shape_id, c_ai.shape_id])

    # =========================================================
    # SLIDE 13: LUYỆN TẬP: TRẮC NGHIỆM TƯƠNG TÁC
    # =========================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, COLOR_ICE_BG)
    add_slide_transition(s13, "push")
    create_header(s13, "HĐ 3 (10p)", "LUYỆN TẬP: TRẮC NGHIỆM CỦNG CỐ")

    # Q1
    q1 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.5), Inches(12.133), Inches(1.7))
    q1.fill.solid()
    q1.fill.fore_color.rgb = COLOR_WHITE
    q1.line.color.rgb = COLOR_ROYAL_BLUE
    q1.line.width = Pt(2.0)
    tf_q1 = q1.text_frame
    p_q1 = tf_q1.paragraphs[0]
    p_q1.text = "CÂU 1: Phương trình x = 6cos(2πt - π/3) cm có biên độ và tần số là:"
    p_q1.font.size = Pt(24)
    p_q1.font.bold = True
    p_q1.font.color.rgb = COLOR_NAVY_DARK

    opt1_a = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.35), Inches(5.8), Inches(0.7))
    opt1_a.fill.solid()
    opt1_a.fill.fore_color.rgb = COLOR_EMERALD_BG
    opt1_a.line.color.rgb = COLOR_EMERALD
    opt1_a.line.width = Pt(2.5)
    opt1_a.text = " A. A = 6 cm ; f = 1 Hz. [ĐÚNG ✔]"
    opt1_a.text_frame.paragraphs[0].font.size = Pt(22)
    opt1_a.text_frame.paragraphs[0].font.bold = True
    opt1_a.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY_DARK
    
    opt1_b = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.35), Inches(5.7), Inches(0.7))
    opt1_b.fill.solid()
    opt1_b.fill.fore_color.rgb = COLOR_WHITE
    opt1_b.line.color.rgb = COLOR_CARD_BORDER
    opt1_b.text = " B. A = 6 cm ; f = 2 Hz."
    opt1_b.text_frame.paragraphs[0].font.size = Pt(22)
    opt1_b.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_MAIN

    # Q2
    q2 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(3.35), Inches(12.133), Inches(1.7))
    q2.fill.solid()
    q2.fill.fore_color.rgb = COLOR_WHITE
    q2.line.color.rgb = COLOR_ROYAL_BLUE
    q2.line.width = Pt(2.0)
    tf_q2 = q2.text_frame
    p_q2 = tf_q2.paragraphs[0]
    p_q2.text = "CÂU 2: Trong 1 chu kì, quãng đường vật dao động điều hòa đi được là:"
    p_q2.font.size = Pt(24)
    p_q2.font.bold = True
    p_q2.font.color.rgb = COLOR_NAVY_DARK

    opt2_a = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(5.8), Inches(0.7))
    opt2_a.fill.solid()
    opt2_a.fill.fore_color.rgb = COLOR_WHITE
    opt2_a.line.color.rgb = COLOR_CARD_BORDER
    opt2_a.text = " A. S = 2A."
    opt2_a.text_frame.paragraphs[0].font.size = Pt(22)
    opt2_a.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_MAIN

    opt2_c = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.2), Inches(5.7), Inches(0.7))
    opt2_c.fill.solid()
    opt2_c.fill.fore_color.rgb = COLOR_EMERALD_BG
    opt2_c.line.color.rgb = COLOR_EMERALD
    opt2_c.line.width = Pt(2.5)
    opt2_c.text = " C. S = 4A. [ĐÚNG ✔]"
    opt2_c.text_frame.paragraphs[0].font.size = Pt(22)
    opt2_c.text_frame.paragraphs[0].font.bold = True
    opt2_c.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY_DARK

    # Q3
    q3 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.2), Inches(12.133), Inches(1.7))
    q3.fill.solid()
    q3.fill.fore_color.rgb = COLOR_WHITE
    q3.line.color.rgb = COLOR_ROYAL_BLUE
    q3.line.width = Pt(2.0)
    tf_q3 = q3.text_frame
    p_q3 = tf_q3.paragraphs[0]
    p_q3.text = "CÂU 3: Đồ thị li độ – thời gian của dao động điều hòa có dạng:"
    p_q3.font.size = Pt(24)
    p_q3.font.bold = True
    p_q3.font.color.rgb = COLOR_NAVY_DARK

    opt3_b = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(5.8), Inches(0.7))
    opt3_b.fill.solid()
    opt3_b.fill.fore_color.rgb = COLOR_EMERALD_BG
    opt3_b.line.color.rgb = COLOR_EMERALD
    opt3_b.line.width = Pt(2.5)
    opt3_b.text = " B. Đường hình sin. [ĐÚNG ✔]"
    opt3_b.text_frame.paragraphs[0].font.size = Pt(22)
    opt3_b.text_frame.paragraphs[0].font.bold = True
    opt3_b.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY_DARK

    opt3_a = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(6.05), Inches(5.7), Inches(0.7))
    opt3_a.fill.solid()
    opt3_a.fill.fore_color.rgb = COLOR_WHITE
    opt3_a.line.color.rgb = COLOR_CARD_BORDER
    opt3_a.text = " A. Đường parabol."
    opt3_a.text_frame.paragraphs[0].font.size = Pt(22)
    opt3_a.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_MAIN

    add_click_animations(s13, [opt1_a.shape_id, opt2_c.shape_id, opt3_b.shape_id])

    # =========================================================
    # SLIDE 14: HOẠT ĐỘNG 4: VẬN DỤNG & NHIỆM VỤ SỐ
    # =========================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, COLOR_ICE_BG)
    add_slide_transition(s14, "push")
    create_header(s14, "HĐ 4 (5p)", "VẬN DỤNG THỰC TIỄN & TỰ HỌC")

    c_v1, top_v1 = add_card_with_autofit(s14, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.3), "ỨNG DỤNG THÁP TAIPEI 101", COLOR_WHITE, COLOR_ROYAL_BLUE, base_title_size=24)
    add_bullet_list(s14, Inches(0.8), top_v1 + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Con lắc cản dịu TMD (Tuned Mass Damper): Quả cầu thép 660 tấn treo tầng 87-92.",
        "Dao động ngược pha với độ rung của tòa nhà khi bão và động đất.",
        "Giảm tới 40% biên độ lắc lư của tòa nhà."
    ], font_size=26, space_after=12)

    c_v2, top_v2 = add_card_with_autofit(s14, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.3), "NHIỆM VỤ TỰ HỌC TẠI NHÀ", COLOR_WHITE, COLOR_EMERALD, base_title_size=24)
    add_bullet_list(s14, Inches(7.0), top_v2 + Inches(0.05), Inches(5.5), Inches(4.3), [
        "Hoàn thành các bài tập trong Phiếu học tập số 1 và số 2.",
        "Dùng AI tra cứu 3 ứng dụng dao động điều hòa trong y tế và công nghệ bán dẫn.",
        "Nộp bài thu hoạch trên hệ thống Padlet/LMS."
    ], font_size=26, space_after=12)

    add_click_animations(s14, [c_v1.shape_id, c_v2.shape_id])

    # =========================================================
    # SLIDE 15: TỔNG KẾT BÀI HỌC
    # =========================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15, COLOR_NAVY_DARK)
    add_slide_transition(s15, "fade")

    card_sum = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(6.3))
    card_sum.fill.solid()
    card_sum.fill.fore_color.rgb = RGBColor(15, 23, 42)
    card_sum.line.color.rgb = COLOR_AMBER
    card_sum.line.width = Pt(3.0)

    title_box_s15 = s15.shapes.add_textbox(Inches(1.0), Inches(0.9), Inches(11.333), Inches(0.8))
    tf_s15 = title_box_s15.text_frame
    p_s15 = tf_s15.paragraphs[0]
    p_s15.text = "TỔNG KẾT BÀI HỌC & ĐIỀU HƯỚNG"
    p_s15.font.size = Pt(34)
    p_s15.font.bold = True
    p_s15.font.color.rgb = COLOR_AMBER
    p_s15.alignment = PP_ALIGN.CENTER
    p_s15.font.name = 'Segoe UI'

    keywords = [
        ("PHƯƠNG TRÌNH LI ĐỘ", "x = A·cos(ωt + φ) với A > 0, ω > 0."),
        ("CHU KÌ & TẦN SỐ", "T = 2π/ω = 1/f ; f = 1/T = ω/2π."),
        ("VÒNG TRÒN LƯỢNG GIÁC", "Hình chiếu chuyển động tròn đều là dao động điều hòa."),
        ("ĐỒ THỊ HÌNH SIN (x-t)", "Trạng thái lặp lại tuần hoàn sau mỗi chu kì T.")
    ]
    for idx, (kw_title, kw_desc) in enumerate(keywords):
        x = Inches(1.2 + (idx % 2) * 5.6)
        y = Inches(1.9 + (idx // 2) * 1.8)
        kw_box = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.3), Inches(1.55))
        kw_box.fill.solid()
        kw_box.fill.fore_color.rgb = COLOR_NAVY_DARK
        kw_box.line.color.rgb = COLOR_ROYAL_BLUE
        kw_box.line.width = Pt(2.0)
        tf_kw = kw_box.text_frame
        tf_kw.word_wrap = True
        p_t = tf_kw.paragraphs[0]
        p_t.text = f"🔑 {kw_title}"
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_AMBER
        p_d = tf_kw.add_paragraph()
        p_d.text = kw_desc
        p_d.font.size = Pt(20)
        p_d.font.color.rgb = COLOR_WHITE

    next_box = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.5), Inches(10.9), Inches(1.1))
    next_box.fill.solid()
    next_box.fill.fore_color.rgb = COLOR_NAVY_LIGHT
    next_box.line.color.rgb = COLOR_EMERALD
    next_box.line.width = Pt(2.0)
    tf_next = next_box.text_frame
    p_n1 = tf_next.paragraphs[0]
    p_n1.text = "📖 ĐỌC TRƯỚC BÀI 2: MÔ TẢ DAO ĐỘNG ĐIỀU HÒA (VẬN TỐC & GIA TỐC)"
    p_n1.font.size = Pt(22)
    p_n1.font.bold = True
    p_n1.font.color.rgb = COLOR_EMERALD
    p_n1.alignment = PP_ALIGN.CENTER

    output_pptx = Path("dau-ra/bai-day/vat-ly-11/bai-1-dao-dong-dieu-hoa/Slide-PowerPoint-Bai-1-Vat-li-11.pptx")
    prs.save(str(output_pptx))
    print(f"[SUCCESS] Đã lưu bài giảng PowerPoint Vật lí 11 Bài 1: {output_pptx}")

if __name__ == "__main__":
    generate_vl11_pptx()
