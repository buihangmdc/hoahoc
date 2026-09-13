import sys
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
except ImportError:
    print("Vui lòng cài đặt python-pptx: py -m pip install python-pptx")
    sys.exit(1)

def markdown_to_pptx(md_path, pptx_path):
    prs = Presentation()
    # Tỉ lệ màn hình rộng 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    slides_content = []
    current_slide = None

    for line in lines:
        line_clean = line.strip()
        if not line_clean:
            continue
        
        # Tiêu đề Slide mới
        if line_clean.startswith('## '):
            if current_slide:
                slides_content.append(current_slide)
            title = line_clean.replace('## ', '').replace('**', '')
            current_slide = {'title': title, 'bullets': []}
        elif line_clean.startswith('# '):
            if current_slide:
                slides_content.append(current_slide)
            title = line_clean.replace('# ', '').replace('**', '')
            current_slide = {'title': title, 'bullets': ['Slide mở đầu']}
        elif current_slide is not None:
            if line_clean.startswith(('-', '*', '+')) or re.match(r'^\d+\.', line_clean):
                bullet_text = re.sub(r'^[-*+]\s+|\d+\.\s+', '', line_clean)
                current_slide['bullets'].append(bullet_text.replace('**', ''))
            elif not line_clean.startswith('---'):
                current_slide['bullets'].append(line_clean.replace('**', ''))

    if current_slide:
        slides_content.append(current_slide)

    blank_slide_layout = prs.slide_layouts[6]

    for slide_data in slides_content:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Tiêu đề slide
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = slide_data['title']
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(16, 44, 87)
        p_title.font.name = 'Arial'

        # Nội dung chi tiết
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for idx, bullet in enumerate(slide_data['bullets']):
            p = tf_content.paragraphs[0] if idx == 0 else tf_content.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(33, 37, 41)
            p.font.name = 'Calibri'
            p.space_after = Pt(10)

    prs.save(pptx_path)
    print(f"Đã xuất bài giảng PowerPoint thành công: {pptx_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Cách dùng: py scripts/export_to_pptx.py <file-input.md> <file-output.pptx>")
        sys.exit(1)
    
    input_file = Path(sys.argv[1])
    output_file = Path(sys.argv[2])
    markdown_to_pptx(input_file, output_file)